"""Export transcripts.jsonl to a formatted Excel workbook.

Sheet "Transkrip": one row per video, caption next to transcript, filters, frozen
header, clickable links. Word count, spoken-mention and speaker-role columns are
formulas.
Sheet "Ringkasan": status counts, speed and speaker roles, as formulas over "Transkrip".

Full captions are re-read from OpenSearch (_mget, read-only) when --netrc-file is
given; otherwise the 300-char caption from the manifest is used. Voiceprint columns
are added when the batch folder has a speakers.jsonl (from speaker-id/speaker_id.py).

Usage:
  python export_xlsx.py --dir <batch folder> --out <file.xlsx> [--netrc-file <path>]
"""
import argparse
import json
import pathlib
import re

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from fetch_videos import DEFAULT_URL, credentials
import requests

STATUS_ORDER = {"done": 0, "suspect_hallucination": 1, "no_speech": 2, "no_audio": 3, "failed": 4}
STATUS_INFO = {
    "done": "Ada ucapan, transkrip terisi",
    "suspect_hallucination": "Whisper mengarang frasa umum (mis. 'Terima kasih.') pada klip tanpa ucapan",
    "no_speech": "Hanya musik / hening",
    "no_audio": "File video tanpa track audio",
    "failed": "Error saat transkripsi",
}
ROLE_INFO = {
    "pembicara": "Suara Prabowo terdeteksi minimal 3 detik; kutipan boleh diatribusikan ke Prabowo",
    "dibicarakan": "Bukan suara Prabowo, tapi ucapan menyebut Prabowo/presiden",
    "hanya_caption": "Prabowo hanya ada di caption/hashtag, tidak di ucapan",
    "tidak_disebut": "Prabowo tidak ada di ucapan maupun caption (mis. video dari link manual)",
}
# (key, header, width, wrap)
COLS = [("no", "No", 5, False), ("status", "Status", 20, False), ("catatan", "Catatan", 28, True),
        ("platform", "Platform", 10, False), ("created_at", "Created at", 18, False),
        ("durasi", "Durasi audio (detik)", 11, False), ("proses", "Waktu proses (detik)", 11, False),
        ("lang", "Bahasa transkrip", 10, False), ("kata", "Jumlah kata", 9, False),
        ("mbg", "Sebut MBG (lisan)", 10, False), ("prabowo_lisan", "Sebut Prabowo/presiden (lisan)", 12, False),
        ("suara", "Suara Prabowo (voiceprint)", 13, False), ("suara_detik", "Durasi suara Prabowo (detik)", 12, False),
        ("suara_porsi", "Porsi suara Prabowo dari durasi video", 12, False),
        ("suara_skor", "Skor kemiripan suara tertinggi", 11, False),
        ("n_suara", "Jumlah suara berbeda (perkiraan)", 11, False), ("peran", "Peran Prabowo", 14, False),
        ("caption", "Caption", 55, True), ("transkrip", "Transkrip", 90, True),
        ("kutipan", "Kutipan suara Prabowo (detik ke-)", 60, True),
        ("link", "Link", 22, False), ("file", "File", 30, False), ("_id", "_id", 22, False), ("_index", "_index", 24, False)]
SPEAKER_KEYS = {"suara", "suara_detik", "suara_porsi", "suara_skor", "n_suara", "peran", "kutipan"}
FONT = "Arial"


def full_captions(recs, netrc_file):
    user, pwd = credentials(DEFAULT_URL, netrc_file)
    # Videos added from a manual link have no _index: they are not in OpenSearch.
    docs = [{"_index": r["_index"], "_id": r["_id"], "_source": ["content"]} for r in recs if r.get("_index")]
    if not docs:
        return {}
    r = requests.post(f"{DEFAULT_URL}/_mget", json={"docs": docs}, auth=(user, pwd), timeout=60)
    r.raise_for_status()
    return {d["_id"]: d["_source"].get("content", "") for d in r.json()["docs"] if d.get("found")}


def notes_for(r, sp, first_seen):
    notes, text = [], r.get("transcript") or ""
    if r["transcript_status"] != "done":
        return ""
    if len(text.split()) <= 3:
        notes.append("Transkrip sangat pendek, cek manual")
    if r.get("transcript_lang") != r.get("lang"):
        notes.append(f"Bahasa beda dari caption ({r.get('lang')})")
    dropped = [s for s in r.get("segments", []) if s.get("hallucination") not in (None, "kosong")]
    if dropped:
        secs = sum(s["end"] - s["start"] for s in dropped)
        reasons = ", ".join(sorted({s["hallucination"] for s in dropped}))
        notes.append(f"{len(dropped)} segmen halusinasi Whisper dibuang ({secs:.0f} detik, {reasons})")
    if r.get("prompt_fallback"):
        notes.append("Prompt ejaan memicu halusinasi, ditranskrip ulang tanpa prompt: ejaan istilah bisa meleset "
                     "(mis. 'MBT' untuk MBG)")
    if sp and sp.get("speech_seconds"):
        short = sum(v["seconds"] for v in sp.get("voices", []) if v["seconds"] < 3)
        if short / sp["speech_seconds"] >= 0.5:
            notes.append(f"{short / sp['speech_seconds']:.0%} ucapan berupa giliran bicara pendek (<3 detik per suara): "
                         "jumlah suara tidak terhitung, bisa lebih banyak")
    key = text[:120]
    if key in first_seen:
        notes.append(f"Duplikat isi dengan _id {first_seen[key]} (reupload)")
    else:
        first_seen[key] = r["_id"]
    return "; ".join(notes)


def mmss(t):
    return f"{int(t // 60)}:{int(t % 60):02d}"


def build(recs, files, captions, manual_notes, speakers, ref_source, summary_notes, out):
    cols = [c for c in COLS if speakers or c[0] not in SPEAKER_KEYS]
    L = {key: get_column_letter(i) for i, (key, *_) in enumerate(cols, 1)}
    idx = {key: i for i, (key, *_) in enumerate(cols, 1)}
    wb = Workbook()
    ws = wb.active
    ws.title = "Transkrip"
    head_font = Font(name=FONT, bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="1F4E78")
    body_font = Font(name=FONT, size=10)
    link_font = Font(name=FONT, size=10, color="0563C1", underline="single")

    for c, (_, name, width, _) in enumerate(cols, 1):
        cell = ws.cell(1, c, name)
        cell.font, cell.fill = head_font, head_fill
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        ws.column_dimensions[get_column_letter(c)].width = width
    ws.row_dimensions[1].height = 45

    first_seen = {}
    for i, r in enumerate(recs, 2):
        note = "; ".join(filter(None, [notes_for(r, speakers.get(r["_id"]), first_seen), manual_notes.get(r["_id"], "")]))
        m, sp = f"{L['transkrip']}{i}", speakers.get(r["_id"])
        v = {
            "no": i - 1, "status": r["transcript_status"], "catatan": note, "platform": r.get("platform"),
            "created_at": r.get("created_at"), "durasi": r.get("transcript_duration"),
            "proses": r.get("process_seconds"), "lang": r.get("transcript_lang"),
            "kata": f'=IF(LEN(TRIM({m}))=0,0,LEN(TRIM({m}))-LEN(SUBSTITUTE(TRIM({m})," ",""))+1)',
            "mbg": f'=IF(OR(ISNUMBER(SEARCH("MBG",{m})),ISNUMBER(SEARCH("makan bergizi",{m}))),"ya","tidak")',
            "prabowo_lisan": f'=IF(OR(ISNUMBER(SEARCH("prabowo",{m})),ISNUMBER(SEARCH("presiden",{m}))),"ya","tidak")',
            "suara": sp["verdict"] if sp else "tanpa_ucapan",
            "suara_detik": sp["target_seconds"] if sp else 0,
            "suara_porsi": f'=IF(N({L["durasi"]}{i})=0,0,{L["suara_detik"]}{i}/{L["durasi"]}{i})' if speakers else None,
            "suara_skor": sp.get("max_sim") if sp else None,
            "n_suara": sp.get("n_voices", 0) if sp else 0,
            "peran": (f'=IF({L["suara"]}{i}="ya","pembicara",'
                      f'IF({L["prabowo_lisan"]}{i}="ya","dibicarakan",'
                      f'IF(ISNUMBER(SEARCH("prabowo",{L["caption"]}{i})),"hanya_caption","tidak_disebut")))')
                     if speakers else None,
            "caption": captions.get(r["_id"], r.get("content") or ""), "transkrip": r.get("transcript") or "",
            "kutipan": "\n".join(f"[{mmss(s['start'])}] {s['text']}" for s in (sp or {}).get("segments", [])
                                 if s["label"] == "target"),
            "link": r.get("link"), "file": files.get(r["_id"]), "_id": r["_id"], "_index": r.get("_index"),
        }
        for c, (key, _, _, wrap) in enumerate(cols, 1):
            cell = ws.cell(i, c, v[key])
            cell.font = body_font
            cell.alignment = Alignment(wrap_text=wrap, vertical="top")
        for key, fmt in (("durasi", "0.0"), ("proses", "0.0"), ("suara_detik", "0.0"), ("suara_porsi", "0%"), ("suara_skor", "0.00")):
            if key in idx:
                ws.cell(i, idx[key]).number_format = fmt
        if r.get("link"):
            ws.cell(i, idx["link"]).hyperlink = r["link"]
            ws.cell(i, idx["link"]).font = link_font
        ws.row_dimensions[i].height = 90 if r["transcript_status"] == "done" else 30

    last = len(recs) + 1
    end_col = get_column_letter(len(cols))
    ws.auto_filter.ref = f"A1:{end_col}{last}"
    ws.freeze_panes = "C2"
    fills = {"done": "E2EFDA", "suspect_hallucination": "FCE4D6", "no_speech": "EDEDED", "no_audio": "EDEDED", "failed": "F8CBAD"}
    for status, color in fills.items():
        ws.conditional_formatting.add(f"A2:{end_col}{last}",
                                      FormulaRule(formula=[f'$B2="{status}"'], fill=PatternFill("solid", fgColor=color)))

    s = wb.create_sheet("Ringkasan")
    s.column_dimensions["A"].width = 44
    s.column_dimensions["B"].width = 14
    s.column_dimensions["C"].width = 80
    rng = lambda key: f"Transkrip!${L[key]}$2:${L[key]}${last}"
    small, bold = Font(name=FONT, size=10), Font(name=FONT, size=14, bold=True)
    row = [1]

    def put(a=None, b=None, c=None, header=False, fmt=None):
        for col, val in enumerate((a, b, c), 1):
            if val is not None:
                s.cell(row[0], col, val).font = small
            if header:
                s.cell(row[0], col).font, s.cell(row[0], col).fill = head_font, head_fill
        if fmt:
            s.cell(row[0], 2).number_format = fmt
        row[0] += 1
        return row[0] - 1

    put("Ringkasan batch")
    s["A1"].font = bold
    row[0] += 1
    put("Status", "Jumlah", "Arti", header=True)
    first = row[0]
    for st in STATUS_ORDER:
        put(st, f'=COUNTIF({rng("status")},"{st}")', STATUS_INFO[st])
    put("Total video", f"=SUM(B{first}:B{row[0] - 1})")
    row[0] += 1
    put("Metrik", "Nilai", "Keterangan", header=True)
    dur = put("Total durasi audio (menit)", f'=SUM({rng("durasi")})/60', "Semua video yang berhasil diproses", fmt="0.0")
    proc = put("Total waktu proses (menit)", f'=SUMIF({rng("durasi")},">0",{rng("proses")})/60',
               "CPU, model large-v3-turbo int8", fmt="0.0")
    put("Kecepatan (x real-time)", f"=IFERROR(B{dur}/B{proc},0)", "Durasi audio / waktu proses", fmt="0.0")
    all_match = all(re.search("prabowo", c, re.I) and re.search("mbg", c, re.I)
                    for c in (captions.get(r["_id"], r.get("content") or "") for r in recs))
    put("Video 'done' yang sebut MBG lisan", f'=COUNTIFS({rng("status")},"done",{rng("mbg")},"ya")',
        "Semua caption memuat 'Prabowo' dan 'MBG'" if all_match else None)
    put("Video 'done' yang sebut Prabowo/presiden lisan", f'=COUNTIFS({rng("status")},"done",{rng("prabowo_lisan")},"ya")')
    if speakers:
        row[0] += 1
        put("Peran Prabowo", "Jumlah", "Arti", header=True)
        first = row[0]
        for role, info in ROLE_INFO.items():
            put(role, f'=COUNTIF({rng("peran")},"{role}")', info)
        put("Total video", f"=SUM(B{first}:B{row[0] - 1})")
        put("Total durasi suara Prabowo (detik)", f'=SUM({rng("suara_detik")})', fmt="0.0")
        th = next(iter(speakers.values()))["threshold"]
        row[0] += 1
        put("Metode voiceprint")
        put("Model", "ECAPA-TDNN", "speechbrain/spkrec-ecapa-voxceleb; jendela 3 detik (hop 1,5 detik) di dalam segmen Whisper")
        put("Cara kerja", None, "Jendela dikelompokkan per suara dalam tiap video; rata-rata embedding tiap suara dibandingkan dengan referensi (cosine)")
        put("Referensi suara", None, ref_source)
        put("Ambang (cosine suara vs referensi)", th, "Satu ambang tanpa zona abu-abu: di atas = Prabowo, di bawah = bukan")
        put("Minimal durasi untuk 'ya' (detik)", 3, "Total durasi segmen dari suara yang lolos ambang 'ya'")
        for note in summary_notes:
            put("Catatan", None, note)
    wb.move_sheet("Ringkasan", offset=-1)
    wb.active = 0
    wb.save(out)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--dir", required=True)
    ap.add_argument("--out", required=True)
    ap.add_argument("--netrc-file")
    ap.add_argument("--note", action="append", default=[], help="manual note as _id=text, repeatable")
    ap.add_argument("--ref-source", default="", help="reference clip name shown in the summary sheet")
    ap.add_argument("--summary-note", action="append", default=[], help="extra note line for the summary sheet, repeatable")
    a = ap.parse_args()
    folder = pathlib.Path(a.dir)
    lines = lambda name: [json.loads(l) for l in (folder / name).read_text(encoding="utf-8").splitlines() if l]
    files = {m["_id"]: m.get("file") for m in lines("manifest.jsonl")}
    recs = sorted(lines("transcripts.jsonl"),
                  key=lambda r: (STATUS_ORDER.get(r["transcript_status"], 9), r.get("created_at") or ""))
    speakers = {s["_id"]: s for s in lines("speakers.jsonl")} if (folder / "speakers.jsonl").exists() else {}
    # Full captions are cached in the batch folder, so re-exports work without OpenSearch access.
    cache = folder / "captions.json"
    captions = json.loads(cache.read_text(encoding="utf-8")) if cache.exists() else {}
    missing = [r for r in recs if r["_id"] not in captions and r.get("_index")]
    if a.netrc_file and missing:
        captions |= full_captions(missing, a.netrc_file)
        cache.write_text(json.dumps(captions, ensure_ascii=False, indent=1), encoding="utf-8")
    manual = dict(n.split("=", 1) for n in a.note)
    build(recs, files, captions, manual, speakers, a.ref_source, a.summary_note, a.out)
    print(f"{len(recs)} rows -> {a.out} (full captions: {len(captions)}, speaker rows: {len(speakers)})")


if __name__ == "__main__":
    main()
